import io
import openpyxl
from django.db import transaction
from apps.authentication.models import User, UserRole
from apps.academics.models import Student, StudentEnrollment, ClassroomParallel

class ExcelStudentImporter:
    """
    Servicio de importación de estudiantes desde archivos Microsoft Excel (.xlsx).
    """

    ID_COLUMN_ALIASES = [
        'matricula', 'código', 'codigo', 'id', 'id_estudiante', 'estudiante_id',
        'ci', 'carnet', 'registro', 'ru'
    ]
    NAME_COLUMN_ALIASES = [
        'nombre', 'nombres', 'nombre_completo', 'apellidos_y_nombres', 'alumno',
        'estudiante', 'nombre completo'
    ]
    EMAIL_COLUMN_ALIASES = [
        'email', 'correo', 'correo_electronico', 'correo electrónico'
    ]

    @classmethod
    def import_from_file(cls, file_obj, parallel: ClassroomParallel) -> dict:
        """
        Lee el archivo Excel e inscribe a los estudiantes en el paralelo indicado.
        """
        if hasattr(file_obj, 'read'):
            wb = openpyxl.load_workbook(file_obj, data_only=True)
        else:
            wb = openpyxl.load_workbook(io.BytesIO(file_obj), data_only=True)

        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))

        if not rows or len(rows) < 2:
            return {
                'success': False,
                'total_rows': 0,
                'imported_count': 0,
                'enrolled_count': 0,
                'skipped_count': 0,
                'errors': ['El archivo Excel está vacío o no contiene filas de datos.'],
            }

        headers = [str(cell).strip().lower() if cell is not None else '' for cell in rows[0]]

        id_col_idx = cls._find_column_index(headers, cls.ID_COLUMN_ALIASES)
        name_col_idx = cls._find_column_index(headers, cls.NAME_COLUMN_ALIASES)
        email_col_idx = cls._find_column_index(headers, cls.EMAIL_COLUMN_ALIASES)

        if id_col_idx is None:
            return {
                'success': False,
                'total_rows': len(rows) - 1,
                'imported_count': 0,
                'enrolled_count': 0,
                'skipped_count': len(rows) - 1,
                'errors': [f'No se encontró la columna de ID/Matrícula. Columnas detectadas: {headers}'],
            }

        if name_col_idx is None:
            return {
                'success': False,
                'total_rows': len(rows) - 1,
                'imported_count': 0,
                'enrolled_count': 0,
                'skipped_count': len(rows) - 1,
                'errors': [f'No se encontró la columna de Nombre Completo. Columnas detectadas: {headers}'],
            }

        total_processed = 0
        imported_count = 0
        enrolled_count = 0
        skipped_count = 0
        errors = []

        with transaction.atomic():
            for row_num, row_data in enumerate(rows[1:], start=2):
                if not any(row_data):
                    continue  # Fila completamente vacía

                total_processed += 1

                raw_id = row_data[id_col_idx]
                raw_name = row_data[name_col_idx]
                raw_email = row_data[email_col_idx] if email_col_idx is not None and len(row_data) > email_col_idx else None

                id_code = str(raw_id).strip() if raw_id is not None else ''
                full_name = str(raw_name).strip() if raw_name is not None else ''
                email = str(raw_email).strip().lower() if raw_email is not None else ''

                if not id_code or not full_name:
                    skipped_count += 1
                    errors.append({
                        'row': row_num,
                        'reason': 'Fila con ID o Nombre incompleto'
                    })
                    continue

                # Crear o vincular cuenta de usuario estudiante
                user_email = email if email else f"{id_code.lower()}@estudiante.edu.bo"
                user, _ = User.objects.get_or_create(
                    student_id_code=id_code,
                    defaults={
                        'email': user_email,
                        'first_name': full_name.split()[0] if full_name else 'Estudiante',
                        'last_name': ' '.join(full_name.split()[1:]) if len(full_name.split()) > 1 else '',
                        'role': UserRole.STUDENT,
                    }
                )

                # Crear o actualizar perfil de Estudiante
                student, created_student = Student.objects.update_or_create(
                    student_id_code=id_code,
                    defaults={
                        'user': user,
                        'full_name': full_name,
                        'email': email or user.email,
                    }
                )
                if created_student:
                    imported_count += 1

                # Inscribir en el paralelo
                enrollment, created_enr = StudentEnrollment.objects.get_or_create(
                    student=student,
                    parallel=parallel
                )
                if created_enr:
                    enrolled_count += 1

        return {
            'success': True,
            'total_rows': total_processed,
            'imported_count': imported_count,
            'enrolled_count': enrolled_count,
            'skipped_count': skipped_count,
            'errors': errors,
        }

    @staticmethod
    def _find_column_index(headers: list[str], aliases: list[str]) -> int | None:
        for idx, header in enumerate(headers):
            if header in aliases or any(alias in header for alias in aliases):
                return idx
        return None
